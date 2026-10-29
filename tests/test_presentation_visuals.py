import os
import tempfile
try:
    import pytest
except ImportError:
    pytest = None
from pptx import Presentation

from digitalagedu.core.llm import slide_kit
from digitalagedu.core.llm.visuals import visual_kit
from digitalagedu.core.llm.presentation_designer import execute_presentation_code_sandbox, SlideDeckRepl


def test_visual_kit_3d_and_charts():
    """Verifies that 3D surface and diagnostic charts render and cache properly."""
    with tempfile.TemporaryDirectory() as temp_dir:
        # Override assets dir for test isolation
        visual_kit.output_dir = temp_dir

        # 1. 3D surface orbit
        vid_path, poster_path = visual_kit.render_3d_surface_orbit(
            surface_type="saddle",
            title="Test Saddle Optimization",
            frames=10,
            fps=12
        )
        assert os.path.exists(vid_path), "3D surface MP4 was not created"
        assert os.path.exists(poster_path), "3D surface poster PNG was not created"
        assert os.path.getsize(vid_path) > 0

        # Caching check
        vid_path2, poster_path2 = visual_kit.render_3d_surface_orbit(
            surface_type="saddle",
            title="Test Saddle Optimization",
            frames=10,
            fps=12
        )
        assert vid_path == vid_path2

        # 2. Diagnostic ROC curve
        roc_path = visual_kit.render_roc_curve(auc_score=0.95, title="Test ROC")
        assert os.path.exists(roc_path)
        assert os.path.getsize(roc_path) > 0


def test_visual_kit_network_flow():
    """Verifies network architecture animation generation."""
    with tempfile.TemporaryDirectory() as temp_dir:
        visual_kit.output_dir = temp_dir
        vid_path, poster_path = visual_kit.render_network_flow(
            stage_names=["Input", "Encoder", "Classifier"],
            frames=8,
            fps=12
        )
        assert os.path.exists(vid_path)
        assert os.path.exists(poster_path)


def test_slide_kit_media_embedding():
    """Verifies that add_video and add_image correctly embed media into a Presentation."""
    with tempfile.TemporaryDirectory() as temp_dir:
        visual_kit.output_dir = temp_dir
        vid_path, poster_path = visual_kit.render_3d_surface_orbit(frames=6, fps=12)
        roc_path = visual_kit.render_roc_curve(auc_score=0.92)

        prs = Presentation()
        slide = slide_kit.create_slide(prs)
        slide_kit.add_header(slide, "TEST", "Visual Toolkit Integration Test")

        # Embed video
        movie = slide_kit.add_video(slide, x=0.8, y=2.0, w=5.5, h=4.5, video_path=vid_path, poster_image_path=poster_path)
        assert movie is not None

        # Embed image
        pic = slide_kit.add_image(slide, x=6.6, y=2.0, w=5.8, h=4.5, image_path=roc_path, caption="ROC Diagnostic")
        assert pic is not None

        output_pptx = os.path.join(temp_dir, "test_presentation.pptx")
        prs.save(output_pptx)
        assert os.path.exists(output_pptx)
        assert os.path.getsize(output_pptx) > 0


def test_execute_presentation_sandbox():
    """Verifies sandboxed subprocess execution with slide_kit and visual_kit."""
    sample_code = """
import pptx
import slide_kit
import visual_kit

def build_presentation(prs, slide_kit, visual_kit, telemetry, solution_code):
    s = slide_kit.create_slide(prs)
    slide_kit.add_header(s, "WEEK 1", "Autonomous Presentation Sandbox")
    slide_kit.add_card(s, 0.8, 2.0, 11.7, 4.5, title="Subprocess Execution", body="Verified end-to-end sandbox.")
"""
    with tempfile.TemporaryDirectory() as temp_dir:
        out_deck = os.path.join(temp_dir, "sandbox_deck.pptx")
        success, msg = execute_presentation_code_sandbox(
            python_code=sample_code,
            output_path=out_deck,
            telemetry={"run_summary": {"accuracy": "91.2%"}},
            solution_code="print('Hello World')",
            timeout=180
        )
        assert success is True, f"Sandbox execution failed: {msg}"
        assert os.path.exists(out_deck)
        assert os.path.getsize(out_deck) > 0


def test_execute_presentation_sandbox_error_handling():
    """Verifies that syntax or execution errors inside the script are caught cleanly without crashing."""
    bad_code = """
def build_presentation(prs, slide_kit, visual_kit, telemetry, solution_code):
    raise ValueError("Intentional script error to test non-fatal handling")
"""
    with tempfile.TemporaryDirectory() as temp_dir:
        out_deck = os.path.join(temp_dir, "bad_deck.pptx")
        success, err = execute_presentation_code_sandbox(
            python_code=bad_code,
            output_path=out_deck,
            telemetry={},
            solution_code="",
            timeout=45
        )
        assert success is False
        assert "Intentional script error" in err


def test_slide_deck_repl():
    """Verifies multi-turn interactive REPL session with persistent prs canvas and visual assets."""
    with tempfile.TemporaryDirectory() as temp_dir:
        out_deck = os.path.join(temp_dir, "repl_deck.pptx")
        telemetry = {
            "run_summary": {"accuracy": "93.4%", "auc_roc": "0.981"},
            "contrastive_samples": {
                "top_success": {"image_path": "sample1.jpg", "ground_truth": "Class_A"},
                "hard_failure": {"image_path": "sample2.jpg", "ground_truth": "Class_B", "predicted_class": "Class_A"}
            }
        }
        solution_code = "import torch\ndef forward(x): return x * 2"

        repl = SlideDeckRepl(
            output_path=out_deck,
            telemetry=telemetry,
            solution_code=solution_code,
            temp_dir=temp_dir
        )

        # Turn 1: Asset generation
        step1_code = """
roc_path = visual_kit.render_roc_curve(auc_score=0.98, title="REPL ROC Diagnostic")
visual_assets["roc"] = roc_path
print(f"Generated asset: {roc_path}")
"""
        ok1, msg1 = repl.execute_snippet(step1_code)
        assert ok1 is True, f"Turn 1 failed: {msg1}"
        assert "roc" in repl.visual_assets
        assert os.path.exists(repl.visual_assets["roc"])

        # Turn 2: Slide 1 construction
        step2_code = """
s1 = slide_kit.create_slide(prs)
slide_kit.add_header(s1, "WEEK 1", "REPL Architecture Overview", "Dynamic multi-turn pipeline")
slide_kit.add_metric_card(s1, 0.8, 2.0, 3.6, 1.8, "ACCURACY", "93.4%", subtext="Phase 1 Pipeline")
slide_kit.add_code_box(s1, 4.8, 2.0, 7.7, 4.5, code=solution_code, title="Model Forward Pass")
print("Slide 1 built successfully")
"""
        ok2, msg2 = repl.execute_snippet(step2_code)
        assert ok2 is True, f"Turn 2 failed: {msg2}"
        assert len(repl.prs.slides) == 1

        # Turn 3: Error resilience check (syntax/runtime error caught without crashing REPL)
        bad_snippet = """
this is invalid python syntax ::: !!!
"""
        ok3, err3 = repl.execute_snippet(bad_snippet)
        assert ok3 is False, "Expected error on bad snippet"
        assert "Execution Error" in err3 or "SyntaxError" in err3
        assert len(repl.prs.slides) == 1, "Slide count should not be corrupted by failed turn"

        # Turn 4: Slide 2 construction embedding asset from Turn 1
        step4_code = """
s2 = slide_kit.create_slide(prs)
slide_kit.add_header(s2, "DIAGNOSTICS", "ROC Performance & Decision Thresholds")
slide_kit.add_image(s2, 0.8, 2.0, 6.0, 4.8, image_path=visual_assets["roc"], caption="Test ROC Curve")
slide_kit.add_card(s2, 7.2, 2.0, 5.3, 4.8, title="Ablation Findings", body="Decision threshold aligns with empirical AUC.")
print("Slide 2 built successfully")
"""
        ok4, msg4 = repl.execute_snippet(step4_code)
        assert ok4 is True, f"Turn 4 failed: {msg4}"
        assert len(repl.prs.slides) == 2

        # Finalize
        final_ok = repl.finalize()
        assert final_ok is True
        assert os.path.exists(out_deck)
        assert os.path.getsize(out_deck) > 0


def test_custom_graphic_generation_and_clamping():
    """Verifies that bespoke matplotlib figures and coordinate clamping work properly."""
    with tempfile.TemporaryDirectory() as temp_dir:
        visual_kit.output_dir = temp_dir

        # 1. Test create_figure and save_figure
        fig, ax = visual_kit.create_figure(width=6.0, height=4.0, title="Custom Concept Diagram")
        ax.plot([1, 2, 3], [4, 5, 6], color=visual_kit.Colors.CYAN, lw=2.5)
        ax.annotate("Peak", xy=(3, 6), color=visual_kit.Colors.EMERALD)
        out_img = visual_kit.save_figure(fig, filename="custom_concept.png")
        assert os.path.exists(out_img)
        assert os.path.getsize(out_img) > 0

        # 2. Test slide bounds clamping and layout helpers
        prs = Presentation()
        s = slide_kit.create_slide(prs)
        slide_kit.add_header(s, "CONCEPT", "Neural Feature Map Slicing")

        # Test add_hero_graphic with coordinate overflow (> 7.5 in height)
        hero = slide_kit.add_hero_graphic(s, out_img, x=0.8, y=1.9, w=7.2, h=8.0)
        assert hero is not None
        # Verify clamped height
        assert hero.top.inches + hero.height.inches <= 7.2

        # Test add_callout_column
        callouts = [("Stage 1", "Raw RGB input data"), ("Stage 2", "Feature extraction")]
        cards = slide_kit.add_callout_column(s, callouts, x=8.5, y=1.9, w=4.0, h=5.0)
        assert len(cards) == 2

        # 3. Test SlideDeckRepl get_deck_summary
        out_deck = os.path.join(temp_dir, "test_deck.pptx")
        repl = SlideDeckRepl(out_deck, {}, "", temp_dir)
        step_code = """
s1 = slide_kit.create_slide(prs)
slide_kit.add_header(s1, "INTRO", "Domain Introduction")
"""
        repl.execute_snippet(step_code)
        deck_summary = repl.get_deck_summary()
        assert "Slide 1: 'INTRO'" in deck_summary or "Slide 1" in deck_summary
        assert "Current Slides in Deck (1)" in deck_summary


def test_high_level_diagram_helpers_and_dark_theme():
    """Verifies that high-level diagram helpers produce valid, aesthetic figures with proper contrast."""
    with tempfile.TemporaryDirectory() as temp_dir:
        visual_kit.output_dir = temp_dir

        # 1. Flowchart
        flow_path = visual_kit.draw_flowchart(
            steps=[
                ("Data Ingestion", "Load and normalize image tensors"),
                ("Feature Extractor", "DINOv2 ViT backbone forward pass"),
                ("Diagnostic Head", "Softmax classification projection")
            ],
            title="Computational Pipeline Workflow",
            subtitle="Stage 1 to Stage 3 Data Transformations"
        )
        assert os.path.exists(flow_path)
        assert os.path.getsize(flow_path) > 1000

        # 2. Architecture
        arch_path = visual_kit.draw_architecture(
            layers=[
                {"name": "Input Batch", "detail": "(B, 3, 224, 224)", "op": "Raw RGB"},
                {"name": "ViT Backbone", "detail": "DINOv2 (768-d)", "op": "Frozen Features"},
                {"name": "Linear Classifier", "detail": "fc1(768 -> 2)", "op": "CrossEntropy"}
            ],
            title="Deep Learning Architecture & Tensor Contracts"
        )
        assert os.path.exists(arch_path)
        assert os.path.getsize(arch_path) > 1000

        # 3. Confusion Matrix (both with None and with custom matrix)
        cm_none_path = visual_kit.draw_confusion_matrix(matrix=None, class_names=["Benign", "Malignant"])
        assert os.path.exists(cm_none_path)
        assert os.path.getsize(cm_none_path) > 1000

        cm_custom_path = visual_kit.draw_confusion_matrix(
            matrix=[[450, 50], [30, 470]],
            class_names=["Benign", "Malignant"],
            title="Empirical Diagnostic Confusion Matrix"
        )
        assert os.path.exists(cm_custom_path)
        assert os.path.getsize(cm_custom_path) > 1000

        # 4. Metric Comparison
        metric_path = visual_kit.draw_metric_comparison(
            metrics={"Accuracy": 92.4, "AUC-ROC": 97.1, "Precision": 89.8, "Recall": 93.5},
            title="Cross-Validation Diagnostic Performance"
        )
        assert os.path.exists(metric_path)
        assert os.path.getsize(metric_path) > 1000

        # 5. Distribution (donut with bright text)
        dist_path = visual_kit.draw_distribution(
            class_counts={"Benign": 1800, "Malignant": 1497},
            title="Dataset Class Balance",
            kind="donut"
        )
        assert os.path.exists(dist_path)
        assert os.path.getsize(dist_path) > 1000

        # 6. SafeTelemetryDict never invents data: missing keys raise with the real keys listed
        from digitalagedu.core.llm.slide_deck_repl import SafeTelemetryDict
        safe_dict = SafeTelemetryDict({"Run_Summary": {"auc_roc": 0.99}})
        assert safe_dict.get("confusion") is None
        assert safe_dict["run_summary"]["AUC_ROC"] == 0.99
        try:
            safe_dict["confusion_matrix"]
            raise AssertionError("expected KeyError")
        except KeyError as e:
            assert "Run_Summary" in str(e)


def test_discretized_concept_presentation_architecture():
    """Verifies that the new concept-discretized architecture produces exactly 1 slide per concept,
    embeds domain context and authentic artifacts, and creates a clean 5-slide deck."""
    from unittest.mock import MagicMock
    from digitalagedu.core.llm.presentation_designer import PresentationDesigner
    from digitalagedu.core.llm.schemas import Module, ProblemStatementSchema
    from digitalagedu.core.llm.presentation_prompts import (
        DiscretizedConceptPlan,
        DiscretizedConcept,
        ConceptVisualSpec,
    )
    from types import SimpleNamespace

    def fenced(code):
        # Plain chat completion whose reply is one python fenced block (slide code never goes through JSON)
        message = SimpleNamespace(content=f"Slide code:\n```python\n{code}\n```")
        return SimpleNamespace(choices=[SimpleNamespace(message=message, finish_reason="stop")])

    module = Module(
        id="foliar_necrosis",
        title="Foliar Lesion Boundary Delineation",
        week=3,
        context="Plant pathology foliar necrosis inspection",
        difficulty="Intermediate"
    )
    problem = ProblemStatementSchema(
        title="Foliar Lesion Boundary Delineation",
        domain_context="Detecting early-stage bacterial leaf blight in agricultural crops",
        problem_statement="Segment damaged tissue and compute boundary IoU against ground-truth pathology annotations",
        learning_objectives=["Understand prompt encoding", "Compute binary mask IoU"],
        suggested_focus="Zero-shot promptable segmentation",
        markdown_overview="# Overview"
    )

    designer = PresentationDesigner(client=None)

    # 1. Test that without client, discretization returns None (strictly no fallbacks)
    assert designer._discretize_concepts(module, problem, "", {}) is None
    assert designer.generate_presentation_deck(module, problem, "", {}, "dummy.pptx") is None

    # 2. Test full end-to-end concept-based generation with mock LLM client
    from digitalagedu.core.llm.presentation_prompts import CONCEPT_SLIDE_LAYOUT_TEMPLATE, EvidenceItem

    solution_code = (
        "def normalize_channels(img):\n    return img / img.max()\n\n"
        "def project_prompts(points):\n    return [p * 2 for p in points]\n\n"
        "def compute_iou(a, b):\n    inter = (a & b).sum()\n    return inter / (a | b).sum()\n"
    )
    telemetry = {"run_summary": {"overall_accuracy_percent": 94.5, "class_balance": {"Apple___scab": 13, "Apple___healthy": 40}}}

    def concept(i, title, function, value, category):
        return DiscretizedConcept(
            concept_index=i, title=title,
            algorithmic_principle=f"{title} principle", domain_application=f"{title} on crop canopies",
            pipeline_data_flow="[B, 3, 224, 224] -> [B, 196, 768]",
            key_takeaways=[f"{title} works because the mechanism holds.", f"{title} fails when the assumption breaks."],
            formula_or_rule=f"rule_{i} = |A ∩ B| / |A ∪ B|",
            real_evidence=[EvidenceItem(label="Accuracy", value=value, context="this run")],
            code_function=function, misconception=f"{title} is not magic.",
            speaker_notes=f"Notes for {title}: the mechanism, the number {value} and the code.",
            visual_spec=ConceptVisualSpec(category=category, visual_idea="bar chart", suggested_recipe_query="bar chart"),
        )

    concept_plan = DiscretizedConceptPlan(
        deck_title="Foliar Lesion Boundary Delineation",
        domain_subtitle="Detecting early-stage bacterial leaf blight in agricultural crops",
        pedagogical_narrative="Progression from raw leaf imaging to promptable mask delineation and IoU validation.",
        concepts=[
            concept(1, "Multispectral Foliar Normalization", "normalize_channels", "94.5%", "flowchart"),
            concept(2, "Promptable Spatial Feature Projection", "project_promts", "13 images", "network_architecture"),  # misspelled: closest match is used
            concept(3, "Intersection-over-Union Diagnostic Benchmarking", "compute_iou", "40 images", "metric_comparison"),
        ],
        synthesis_takeaways=["Normalization stabilizes features.", "IoU quantifies lesion overlap."],
    )

    real_chart = ('balance = telemetry["run_summary"]["class_balance"]\n'
                  'ax.bar(list(balance), list(balance.values()))')

    def template_code(i):
        # The prompt's layout template with the chart part filled from real telemetry
        return CONCEPT_SLIDE_LAYOUT_TEMPLATE.replace("{index}", str(i)).replace(
            "# ... plot the real data for this concept here ...", real_chart)

    fabricated = fenced("y = np.random.choice(['Healthy', 'Blight'], size=100)\n" + template_code(1))
    no_code_box = fenced(template_code(2).replace("code=solution_excerpt", "code='pass'").replace(
        'title=code_function_name + "()  - verified solution"', 'title="snippet"'))

    mock_client = MagicMock()
    mock_client.chat.completions.create.side_effect = [concept_plan]
    code_client = MagicMock()
    code_client.chat.completions.create.side_effect = [
        fabricated, fenced(template_code(1)),   # concept 1: fake data rejected unrun, then real data
        no_code_box, fenced(template_code(2)),  # concept 2: verified code missing from the slide, then fixed
        fenced(template_code(3)),
    ]

    with tempfile.TemporaryDirectory() as temp_dir:
        out_deck = os.path.join(temp_dir, "concept_architecture_deck.pptx")
        from digitalagedu.core.llm.generation_trace import GenerationTrace
        trace = GenerationTrace(temp_dir, module.id, "qwen")
        active_designer = PresentationDesigner(client=mock_client, code_client=code_client)
        result_path = active_designer.generate_presentation_deck(
            module=module,
            problem_formulation=problem,
            solution_code=solution_code,
            telemetry=telemetry,
            output_path=out_deck,
            artifacts={"raw_sample": "sample.jpg"},
            trace=trace
        )
        slide_execs = [a for a in trace.doc["activity"].values() if a["prov:type"] == "dage:slide_exec"]
        assert [(a["dage:success"], a["dage:slide_added"]) for a in slide_execs] == [
            (False, False), (True, True), (False, False), (True, True), (True, True)
        ]
        assert "fabricates data" in slide_execs[0]["dage:error"] and "missing required content" in slide_execs[2]["dage:error"]
        assert result_path == out_deck
        assert os.path.exists(out_deck)
        prs = Presentation(out_deck)
        assert len(prs.slides) == 5, f"Expected exactly 5 slides (Title + 3 Concepts + Synthesis), got {len(prs.slides)}"

        concept_slide = prs.slides[2]
        text = " ".join(sh.text_frame.text for sh in concept_slide.shapes if sh.has_text_frame)
        assert "project_prompts" in text and "13 images" in text and "rule_2" in text  # code box, evidence, formula
        assert "Key Point" not in text and "Placeholder" not in text
        assert any(sh.shape_type == 13 for sh in concept_slide.shapes)  # the chart saved via a relative plt.savefig name
        assert "Notes for Promptable Spatial Feature Projection" in concept_slide.notes_slide.notes_text_frame.text

    # 3. A slide that still fails after its self-healing retry crashes the deck (fail fast, no partial decks)
    bad_code = fenced("raise RuntimeError('Forced test error')")
    failing_client, failing_code_client = MagicMock(), MagicMock()
    failing_client.chat.completions.create.side_effect = [concept_plan]
    failing_code_client.chat.completions.create.side_effect = [bad_code, bad_code]
    with tempfile.TemporaryDirectory() as temp_dir:
        out_deck2 = os.path.join(temp_dir, "failed_deck.pptx")
        with pytest.raises(RuntimeError, match="after self-healing retry"):
            PresentationDesigner(client=failing_client, code_client=failing_code_client).generate_presentation_deck(
                module=module,
                problem_formulation=problem,
                solution_code="# verified reference solution code",
                telemetry={},
                output_path=out_deck2
            )
        assert not os.path.exists(out_deck2)


if __name__ == "__main__":
    print("Running test_visual_kit_3d_and_charts...")
    test_visual_kit_3d_and_charts()
    print("Running test_visual_kit_network_flow...")
    test_visual_kit_network_flow()
    print("Running test_slide_kit_media_embedding...")
    test_slide_kit_media_embedding()
    print("Running test_execute_presentation_sandbox...")
    test_execute_presentation_sandbox()
    print("Running test_execute_presentation_sandbox_error_handling...")
    test_execute_presentation_sandbox_error_handling()
    print("Running test_slide_deck_repl...")
    test_slide_deck_repl()
    print("Running test_custom_graphic_generation_and_clamping...")
    test_custom_graphic_generation_and_clamping()
    print("Running test_high_level_diagram_helpers_and_dark_theme...")
    test_high_level_diagram_helpers_and_dark_theme()
    print("Running test_discretized_concept_presentation_architecture...")
    test_discretized_concept_presentation_architecture()
    print("\nALL PRESENTATION VISUAL, REPL & DESIGNER TESTS PASSED SUCCESSFULLY!")





def test_missing_image_fails_the_snippet_and_leaves_no_half_built_slide():
    """add_image used to draw an '[Asset Placeholder: ...]' card that shipped as slide content."""
    from digitalagedu.core.llm.slide_deck_repl import SlideDeckRepl
    with tempfile.TemporaryDirectory() as temp_dir:
        repl = SlideDeckRepl(os.path.join(temp_dir, "deck.pptx"), {}, "", temp_dir)
        ok, msg = repl.execute_snippet(
            "slide_kit.add_split_slide(prs, title='T', left_content=['a'], image_path='missing_chart.png')"
        )
        assert not ok and "image not found: 'missing_chart.png'" in msg
        assert len(repl.prs.slides) == 0  # the slide created before the error was rolled back
        ok, _ = repl.execute_snippet("slide_kit.add_split_slide(prs, title='T', left_content=['a'])")
        assert ok and len(repl.prs.slides) == 1 and repl.finalize()
