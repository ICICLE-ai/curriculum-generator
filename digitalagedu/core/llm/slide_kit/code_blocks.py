from typing import List, Optional, Union
from pptx.util import Pt, Length
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

from digitalagedu.core.llm.slide_kit.theme import (
    Theme,
    _to_length,
    _clamp_bounds,
)


def add_code_box(
    slide,
    x: Union[int, float, Length],
    y: Union[int, float, Length],
    w: Union[int, float, Length],
    h: Union[int, float, Length],
    code: Optional[str] = None,
    title: Optional[str] = None,
    font_size: float = 9.5,
    *args,
    **kwargs
):
    """Renders a syntax-styled dark IDE code block with monospaced font."""
    raw_code = code or kwargs.get("code_string") or kwargs.get("code_str") or kwargs.get("text") or ""
    t = title or kwargs.get("heading") or kwargs.get("header")

    lx, ly, lw, lh = _clamp_bounds(x, y, w, h)

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        lx, ly, lw, lh
    )
    card.fill.solid()
    card.fill.fore_color.rgb = Theme.CODE_BG
    card.line.color.rgb = Theme.CARD_BORDER
    card.line.width = Pt(1.5)

    pad_x = _to_length(0.2)
    pad_y = _to_length(0.15)
    tbox = slide.shapes.add_textbox(
        lx + pad_x, ly + pad_y,
        max(_to_length(0.5), lw - (pad_x * 2)),
        max(_to_length(0.5), lh - (pad_y * 2))
    )
    tf = tbox.text_frame
    tf.word_wrap = True

    first = True
    if t:
        p_t = tf.paragraphs[0]
        p_t.text = str(t).upper()
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = Theme.ACCENT_CYAN
        first = False

    lines = [l for l in str(raw_code).strip().split("\n") if not l.startswith('"""') and not l.startswith("'''")]
    display_code = "\n".join(lines[:24]) if lines else str(raw_code)

    p_c = tf.paragraphs[0] if first else tf.add_paragraph()
    p_c.text = display_code
    p_c.font.name = "Consolas"
    p_c.font.size = Pt(font_size)
    p_c.font.color.rgb = Theme.CODE_TEXT
    if not first:
        p_c.space_before = Pt(4)

    return card


def add_table(
    slide,
    x: Union[int, float, Length],
    y: Union[int, float, Length],
    w: Union[int, float, Length],
    h: Union[int, float, Length],
    headers: Optional[List[str]] = None,
    rows: Optional[List[List[str]]] = None,
    *args,
    **kwargs
):
    """Renders a styled data comparison table on the slide."""
    hdrs = headers or kwargs.get("columns") or []
    rws = rows or kwargs.get("data") or []
    num_rows = len(rws) + 1
    num_cols = len(hdrs)

    lx, ly, lw, lh = _to_length(x), _to_length(y), _to_length(w), _to_length(h)

    table_shape = slide.shapes.add_table(
        num_rows, num_cols, lx, ly, lw, lh
    )
    table = table_shape.table

    for c_idx, head in enumerate(hdrs):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = Theme.CARD_BG
        p = cell.text_frame.paragraphs[0]
        p.text = str(head).upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = Theme.ACCENT_CYAN
        p.alignment = PP_ALIGN.CENTER

    for r_idx, row in enumerate(rws):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = Theme.CODE_BG
            p = cell.text_frame.paragraphs[0]
            p.text = str(val)
            p.font.size = Pt(10)
            p.font.color.rgb = Theme.TEXT_PRIMARY
            p.alignment = PP_ALIGN.CENTER

    return table_shape

