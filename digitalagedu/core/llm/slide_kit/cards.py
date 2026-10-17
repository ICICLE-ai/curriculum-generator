import os
from typing import Dict, Any, List, Optional, Tuple, Union
from pptx.util import Pt, Length
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

from digitalagedu.core.llm.slide_kit.theme import (
    Theme,
    _to_length,
    _clamp_bounds,
)


def add_card(
    slide,
    x: Union[int, float, Length],
    y: Union[int, float, Length],
    w: Union[int, float, Length],
    h: Union[int, float, Length],
    title: Optional[str] = None,
    body: Optional[str] = None,
    bg_color: Optional[RGBColor] = None,
    border_color: Optional[RGBColor] = None,
    accent_color: Optional[RGBColor] = None,
    title_size: int = 12,
    body_size: int = 12,
    *args,
    **kwargs
):
    """Adds a rounded rectangular card container with optional title and body paragraphs."""
    t = title or kwargs.get("heading") or kwargs.get("header")
    b = body or kwargs.get("text") or kwargs.get("content") or kwargs.get("description")
    bg = bg_color or kwargs.get("bg") or Theme.CARD_BG
    border = border_color or kwargs.get("border") or Theme.CARD_BORDER
    acc = accent_color or kwargs.get("accent") or Theme.ACCENT_CYAN

    lx, ly, lw, lh = _clamp_bounds(x, y, w, h)

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        lx, ly, lw, lh
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg
    card.line.color.rgb = border
    card.line.width = Pt(1.5)

    if t or b:
        pad_x = _to_length(0.2)
        pad_y = _to_length(0.18)
        tbox = slide.shapes.add_textbox(
            lx + pad_x, ly + pad_y,
            max(_to_length(0.5), lw - (pad_x * 2)),
            max(_to_length(0.5), lh - (pad_y * 2))
        )
        tf = tbox.text_frame
        tf.word_wrap = True

        first = True
        if t:
            p_t = tf.paragraphs[0]
            p_t.text = str(t)
            p_t.font.size = Pt(title_size)
            p_t.font.bold = True
            p_t.font.color.rgb = acc
            first = False

        if b:
            p_b = tf.paragraphs[0] if first else tf.add_paragraph()
            p_b.text = str(b)
            p_b.font.size = Pt(body_size)
            p_b.font.color.rgb = Theme.TEXT_PRIMARY
            if not first:
                p_b.space_before = Pt(6)

    return card


def add_metric_card(
    slide,
    x: Union[int, float, Length],
    y: Union[int, float, Length],
    w: Union[int, float, Length],
    h: Union[int, float, Length],
    label: Optional[str] = None,
    value: Optional[Any] = None,
    subtext: Optional[str] = None,
    accent_color: Optional[RGBColor] = None,
    *args,
    **kwargs
):
    """Renders a high-impact KPI statistic card."""
    lbl = label or kwargs.get("title") or kwargs.get("name") or ""
    val = value if value is not None else (kwargs.get("val") or kwargs.get("metric") or "")
    sub = subtext or kwargs.get("subtitle") or kwargs.get("sub") or kwargs.get("description")
    acc = accent_color or kwargs.get("accent") or Theme.ACCENT_CYAN

    lx, ly, lw, lh = _clamp_bounds(x, y, w, h)

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        lx, ly, lw, lh
    )
    card.fill.solid()
    card.fill.fore_color.rgb = Theme.CARD_BG
    card.line.color.rgb = Theme.CARD_BORDER
    card.line.width = Pt(1.5)

    pad_x = _to_length(0.15)
    pad_y = _to_length(0.12)
    tbox = slide.shapes.add_textbox(
        lx + pad_x, ly + pad_y,
        max(_to_length(0.5), lw - (pad_x * 2)),
        max(_to_length(0.5), lh - (pad_y * 2))
    )
    tf = tbox.text_frame
    tf.word_wrap = True

    # Label
    p_lbl = tf.paragraphs[0]
    p_lbl.text = str(lbl).upper()
    p_lbl.font.size = Pt(9)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = Theme.TEXT_MUTED

    # Value
    p_val = tf.add_paragraph()
    p_val.text = str(val)
    p_val.font.size = Pt(20)
    p_val.font.bold = True
    p_val.font.color.rgb = acc
    p_val.space_before = Pt(2)

    # Subtext
    if sub:
        p_sub = tf.add_paragraph()
        p_sub.text = str(sub)
        p_sub.font.size = Pt(9.5)
        p_sub.font.color.rgb = Theme.TEXT_MUTED
        p_sub.space_before = Pt(2)

    return card


def add_contrastive_cards(
    slide,
    x: Union[int, float, Length],
    y: Union[int, float, Length],
    w: Union[int, float, Length],
    h: Union[int, float, Length],
    success_data: Optional[Dict[str, Any]] = None,
    failure_data: Optional[Dict[str, Any]] = None,
    *args,
    **kwargs
):
    """
    Renders side-by-side diagnostic case studies (Top Success vs Hard Failure).
    Accepts any keyword alias (success_data, success_dict, success, top_success, failure_data, failure_dict, failure, hard_failure).
    """
    succ = (
        success_data or 
        kwargs.get("success_dict") or 
        kwargs.get("success") or 
        kwargs.get("top_success") or 
        (args[0] if len(args) > 0 else {}) or 
        {}
    )
    fail = (
        failure_data or 
        kwargs.get("failure_dict") or 
        kwargs.get("failure") or 
        kwargs.get("hard_failure") or 
        (args[1] if len(args) > 1 else {}) or 
        {}
    )

    lx, ly, lw, lh = _to_length(x), _to_length(y), _to_length(w), _to_length(h)
    gap = _to_length(0.3)
    col_w = (lw - gap) / 2.0

    # 1. Success Card (Emerald)
    add_card(
        slide,
        x=lx, y=ly, w=col_w, h=lh,
        title="HIGH-CONFIDENCE MATCH (SUCCESS)",
        body=(
            f"Sample: {os.path.basename(succ.get('image_path', succ.get('Sample Path', 'sample.jpg'))) if isinstance(succ, dict) else 'sample.jpg'}\n"
            f"Ground Truth: {succ.get('ground_truth', succ.get('True Category', 'Target')) if isinstance(succ, dict) else 'Target'}\n"
            f"Prediction: {succ.get('predicted_class', succ.get('ground_truth', 'Target')) if isinstance(succ, dict) else 'Target'}\n"
            f"Probabilities: {succ.get('probabilities', 'High Confidence') if isinstance(succ, dict) else 'High Confidence'}"
        ),
        accent_color=Theme.ACCENT_EMERALD,
        title_size=11,
        body_size=11
    )

    # 2. Failure Card (Coral Red)
    add_card(
        slide,
        x=lx + col_w + gap, y=ly, w=col_w, h=lh,
        title="DIAGNOSTIC FAILURE MODE (EDGE CASE)",
        body=(
            f"Sample: {os.path.basename(fail.get('image_path', fail.get('Sample Path', 'failure.jpg'))) if isinstance(fail, dict) else 'failure.jpg'}\n"
            f"Ground Truth: {fail.get('ground_truth', fail.get('True Category', 'True Class')) if isinstance(fail, dict) else 'True Class'}\n"
            f"Predicted: {fail.get('predicted_class', 'Misclassified') if isinstance(fail, dict) else 'Misclassified'}\n"
            f"Probabilities: {fail.get('probabilities', 'Shifted Decision Boundary') if isinstance(fail, dict) else 'Shifted Boundary'}"
        ),
        accent_color=Theme.ACCENT_CORAL,
        title_size=11,
        body_size=11
    )


def add_callout_banner(
    slide,
    x: Union[int, float, Length],
    y: Union[int, float, Length],
    w: Union[int, float, Length],
    h: Union[int, float, Length],
    text: Optional[str] = None,
    title: str = "KEY PEDAGOGICAL TAKEAWAY",
    accent_color: Optional[RGBColor] = None,
    *args,
    **kwargs
):
    """Renders a full-width takeaway callout card with accent header."""
    txt = text or kwargs.get("body") or kwargs.get("content") or ""
    t = title or kwargs.get("heading") or "KEY PEDAGOGICAL TAKEAWAY"
    acc = accent_color or kwargs.get("accent") or Theme.ACCENT_GOLD
    return add_card(
        slide,
        x=x, y=y, w=w, h=h,
        title=str(t),
        body=str(txt),
        accent_color=acc,
        title_size=10,
        body_size=11.5
    )


def add_text_box(slide, x, y, w, h, text, title=None, *args, **kwargs):
    """Alias mapping add_text_box to add_card."""
    return add_card(slide, x, y, w, h, title=title, body=text, *args, **kwargs)

add_text = add_text_box

